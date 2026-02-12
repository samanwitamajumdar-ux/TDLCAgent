from typing_extensions import TypedDict
from langgraph.graph import StateGraph, START, END
from langchain_community.document_loaders import Docx2txtLoader
from agents.document_analyzer_agent import create_doc_analyzer_agent
from agents.requirement_analyzer_agent import create_transcript_analyzer_agent
from agents.epic_story_agent import create_epic_story_creator_agent
from docx import Document
from docx.shared import Pt
from collections import defaultdict
import json
from dotenv import load_dotenv

load_dotenv()

# Graph state
class State(TypedDict):
    feature_list: dict
    requirement_list: dict
    epic_story: dict
    
#Nodes
def generate_features(state: State):
    
    # Load Knowledge Document
    loader = Docx2txtLoader("./example_data/Confluence Design Document.docx")

    data = loader.load()
    
    USER_PROMPT = f"""
    Analyze the following product documentation and extract the canonical feature list.

    DOCUMENTATION:
    {data}
    """

    messages = {"messages":[
        {"role": "user", "content": USER_PROMPT}]
    }

    print("Extracting Features from knowledge source...")
    agent = create_doc_analyzer_agent()
    agent_message = agent.invoke(messages)
    file_path_epic = "./output/features.json"
    with open(file_path_epic, "w", encoding="utf-8") as f:
        json.dump(agent_message["structured_response"], f, indent=2, ensure_ascii=False)
    return {"feature_list": agent_message["structured_response"]}

def generate_requirements(state: State):
    
    # Load Knowledge Document
    loader = Docx2txtLoader("./example_data/Meeting Transcripts.docx")

    data = loader.load()
    
    USER_PROMPT = f"""
    Analyze the following transcript and extract the canonical requirement list.

    TRANSCRIPT:
    {data}
    """

    messages = {"messages":[
        {"role": "user", "content": USER_PROMPT}]
    }

    print("Extracting Requirements from Transcript...")
    agent = create_transcript_analyzer_agent()
    agent_message = agent.invoke(messages)
    file_path_epic = "./output/requirements.json"
    with open(file_path_epic, "w", encoding="utf-8") as f:
        json.dump(agent_message["structured_response"], f, indent=2, ensure_ascii=False)
    return {"requirement_list": agent_message["structured_response"]}

def generate_epic_story(state: State):
    
    USER_PROMPT = f"""
    Analyze the following feature list and requirement list to generate epics and related stories.

    FEATURES:
    {state['feature_list']['features']}
    REQUIREMENTS:
    {state['requirement_list']['transcript_signals']}
    """
    messages = {"messages":[
        {"role": "user", "content": USER_PROMPT}]
    }

    print("Creating the Epics and Stories...")
    agent = create_epic_story_creator_agent()
    agent_message = agent.invoke(messages)
    return {"epic_story": agent_message["structured_response"]}

def export_epics_stories_to_docx(state: State):
    """
    Converts epic/story JSON output into a structured DOCX document.
    Epics and their related stories are grouped together.
    """
    print("Saving Data...")
    data = state["epic_story"]
    
    document = Document()

    # -------------------------------
    # Helper: paragraph formatting
    # -------------------------------
    def add_label_value(label: str, value: str):
        p = document.add_paragraph()
        run_label = p.add_run(f"{label}: ")
        run_label.bold = True
        p.add_run(value if value else "N/A")

    # -------------------------------
    # Index stories by epic_id
    # -------------------------------
    stories_by_epic = defaultdict(list)
    for story in data.get("stories", []):
        stories_by_epic[story["epic_id"]].append(story)

    # -------------------------------
    # Epics & Stories
    # -------------------------------
    for epic in data.get("epics", []):
        # Epic Header
        document.add_heading(
            f"{epic['epic_id']} - {epic['name']}",
            level=1
        )

        add_label_value("Description", epic.get("description"))
        add_label_value("In Scope", epic.get("in_scope"))
        add_label_value("Out of Scope", epic.get("out_of_scope"))

        # Spacer
        document.add_paragraph("")

        # Stories under Epic
        for story in stories_by_epic.get(epic["epic_id"], []):
            document.add_heading(
                f"{story['story_id']} - {story['title']}",
                level=2
            )

            add_label_value("User Story", story.get("user_story"))
            add_label_value("Description", story.get("description"))
            add_label_value("Story Type", story.get("story_type"))

            # Acceptance Criteria
            p = document.add_paragraph()
            run = p.add_run("Acceptance Criteria:")
            run.bold = True
            document.add_paragraph(story.get("acceptance_criteria", "N/A"))

            # Tasks
            p = document.add_paragraph()
            run = p.add_run("Tasks:")
            run.bold = True

            tasks = story.get("tasks", "")
            if tasks:
                for task in [t.strip() for t in tasks.split(",")]:
                    document.add_paragraph(task, style="List Bullet")
            else:
                document.add_paragraph("N/A")

            # Tags
            if story.get("feature_tags"):
                add_label_value("Feature Tags", story.get("feature_tags"))

            if story.get("system_tags"):
                add_label_value("System Tags", story.get("system_tags"))

            # Traceability
            add_label_value("Traceability", story.get("traceability"))

            # Spacer between stories
            document.add_paragraph("")

        # Page break after each epic
        document.add_page_break()

    # -------------------------------
    # Reasoning Log (Audit Section)
    # -------------------------------
    if data.get("reasoning_log"):
        document.add_heading("Reasoning Log (Audit Trail)", level=1)

        for entry in data["reasoning_log"]:
            document.add_heading(
                f"{entry['artifact_type']}: {entry['artifact_name']}",
                level=2
            )
            document.add_paragraph(entry.get("reasoning", "N/A"))

    # -------------------------------
    # Save file
    # -------------------------------
    document.save("./output/Epic_Story.docx")
    
    file_path_epic = "./output/epic_story.json"
    with open(file_path_epic, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=2, ensure_ascii=False)



# Build workflow
parallel_builder = StateGraph(State)

# Add nodes
parallel_builder.add_node("generate_features", generate_features)
parallel_builder.add_node("generate_requirements", generate_requirements)
parallel_builder.add_node("generate_epic_story", generate_epic_story)
parallel_builder.add_node("export_epics_stories_to_docx", export_epics_stories_to_docx)

# Add edges to connect nodes
parallel_builder.add_edge(START, "generate_features")
parallel_builder.add_edge(START, "generate_requirements")
parallel_builder.add_edge("generate_features", "generate_epic_story")
parallel_builder.add_edge("generate_requirements", "generate_epic_story")
parallel_builder.add_edge("generate_epic_story", "export_epics_stories_to_docx")
parallel_builder.add_edge("export_epics_stories_to_docx", END)
parallel_workflow = parallel_builder.compile()

# Invoke
state = parallel_workflow.invoke({})
print("Process Completed.")